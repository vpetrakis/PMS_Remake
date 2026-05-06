"""
TEMPORAL PIPELINE v3.0 — The Truth Engine
MV ALEXIS | MTS Marine Ltd. | Final Production Release
Architecture: Bipartite Veto Engine + 3-Vector Time Reconstruction
"""

import streamlit as st
import pandas as pd
import openpyxl
import io
import re
from datetime import datetime, date
from docx import Document
from rapidfuzz import fuzz

# ─────────────────────────────────────────────────────────────────────────────
# CONFIGURATION & DICTIONARIES
# ─────────────────────────────────────────────────────────────────────────────

st.set_page_config(
    page_title="Temporal Pipeline 3.0",
    page_icon="⚓",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Maritime abbreviation normalization dictionary
MARITIME_ABBREV = {
    r'\bM/?E\b': 'MAIN ENGINE',
    r'\bD/?G\b': 'DIESEL GENERATOR',
    r'\bA/?E\b': 'AUX ENGINE',
    r'\bF/?O\b': 'FUEL OIL',
    r'\bL/?O\b': 'LUBE OIL',
    r'\bH/?F/?O\b': 'HEAVY FUEL OIL',
    r'\bM/?D/?O\b': 'MARINE DIESEL OIL',
    r'\bF/?W\b': 'FRESH WATER',
    r'\bS/?W\b': 'SEA WATER',
    r'\bT/?C\b': 'TURBOCHARGER',
    r'\bC/?W\b': 'COOLING WATER',
    r'\bA/?C\b': 'AIR CONDITIONING',
    r'\bEXH\b': 'EXHAUST',
    r'\bINJ\b': 'INJECTOR',
    r'\bCYL\b': 'CYLINDER',
    r'\bASS\'?Y\b': 'ASSEMBLY',
    r'\bELECT\b': 'ELECTRIC',
    r'\bMOTOT?\b': 'MOTOR',
    r'\bPURIF\b': 'PURIFIER',
    r'\bPRESSURE\b': 'PRESSURE',
    r'\bPUMP\b': 'PUMP',
    r'\bFILTER\b': 'FILTER',
    r'\bVALVE\b': 'VALVE',
    r'\bCOOLER\b': 'COOLER',
}

SYSTEM_TAG_MAP = {
    'ME':  'MAIN ENGINE',
    'DG':  'DIESEL GENERATOR',
    'AE':  'AUX ENGINE',
    'EG':  'EMERGENCY GENERATOR',
    'BL':  'BOILER',
    'EGE': 'BOILER',
    'BG':  'BILGE',
    'FW':  'FRESH WATER',
    'SW':  'SEA WATER',
    'FO':  'FUEL OIL',
    'LO':  'LUBE OIL',
    'HV':  'HVAC',
    'CO':  'COMPRESSOR',
    'ST':  'STEERING',
    'CP':  'CARGO PUMP',
    'BP':  'BALLAST PUMP',
    'PC':  'PUMP',
    'PG':  'PUMP',
    'PH':  'PUMP',
    'FF':  'FAN',
}

# 🛑 CONTEXTUAL QUARANTINE (The Veto Box)
FORBIDDEN_KEYWORDS = {
    'MAIN ENGINE':      ['D/G', 'D.G.', 'DIESEL GENERATOR', 'GENERATOR', 'PURIFIER', 'DUCK KEEL', 'HOLD', 'DECK', 'BOILER'],
    'DIESEL GENERATOR': ['M/E', 'M.E.', 'MAIN ENGINE', 'MAIN ENG', 'BOILER', 'STERN TUBE', 'DUCK KEEL', 'HOLD', 'DECK'],
    'BOILER':           ['M/E', 'MAIN ENGINE', 'D/G', 'DIESEL GENERATOR', 'PURIFIER', 'COMPRESSOR'],
    'COMPRESSOR':       ['M/E', 'D/G', 'BOILER', 'PURIFIER', 'PUMP'],
}

SYSTEM_KEYWORDS = {
    'MAIN ENGINE':        ['M/E', 'M.E.', 'MAIN ENGINE', 'MAIN ENG', 'M.E'],
    'DIESEL GENERATOR':   ['D/G', 'D.G.', 'DIESEL GENERATOR', 'DIESEL GEN', 'GENERATOR ENGINE', 'GEN ENGINE', 'D.G'],
    'BOILER':             ['BOILER', 'BLR', 'EGE', 'ECONOMIZER'],
    'COMPRESSOR':         ['COMPRESSOR', 'COMP.'],
    'PUMP':               ['PUMP'],
    'FAN':                ['FAN', 'EXH. FAN', 'EXH FAN']
}

ACTION_WORDS = {
    'OVERHAUL', 'OVERHAULED', 'REPLACE', 'REPLACED', 'RENEW', 'RENEWED',
    'PULL OUT', 'PULLED OUT', 'DISMANTLE', 'DISMANTLED', 'RECONDITION',
    'RECONDITIONED', 'CHANGE', 'CHANGED', 'FABRICATE', 'FABRICATED',
    'REBUILD', 'REBUILT', 'REPAIR', 'REPAIRED', 'INSTALL', 'INSTALLED',
}

# ─────────────────────────────────────────────────────────────────────────────
# TEXT ANALYSIS HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def extract_system_from_code(code: str) -> str | None:
    if not code or '-' not in code:
        return None
    return SYSTEM_TAG_MAP.get(code.split('-')[0].upper())

def extract_cylinder_from_code(code: str) -> int | None:
    parts = code.split('-')
    if len(parts) >= 2:
        try:
            return int(parts[1])
        except ValueError:
            return None
    return None

def extract_system_from_text(text: str) -> str | None:
    upper = text.upper()
    for system, keywords in SYSTEM_KEYWORDS.items():
        if any(re.search(rf'\b{re.escape(kw)}\b', upper) for kw in keywords):
            return system
    return None

def extract_cylinder_from_text(text: str) -> int | None:
    patterns = [
        r'CYL(?:INDER)?\.?\s*(?:NO\.?|#)?\s*(\d+)',
        r'UNIT\s*(?:NO\.?|#)?\s*(\d+)',
        r'D/?G\s*(?:NO\.?|#)?\s*(\d+)',
        r'A/?E\s*(?:NO\.?|#)?\s*(\d+)',
        r'COMPRESSOR\s*(?:NO\.?|#)?\s*(\d+)',
        r'PUMP\s*(?:NO\.?|#)?\s*(\d+)',
        r'\bNO\.?\s*(\d+)\b',
        r'#\s*(\d+)'
    ]
    upper = text.upper()
    for pat in patterns:
        m = re.search(pat, upper)
        if m:
            return int(m.group(1))
    return None

def normalize_maritime_text(text: str) -> str:
    text = str(text).upper().strip()
    for pattern, replacement in MARITIME_ABBREV.items():
        text = re.sub(pattern, replacement, text, flags=re.IGNORECASE)
    text = re.sub(r'[^\w\s]', ' ', text)
    return re.sub(r'\s+', ' ', text).strip()

def is_action_entry(text: str) -> bool:
    upper = str(text).upper()
    return any(word in upper for word in ACTION_WORDS)

# ─────────────────────────────────────────────────────────────────────────────
# FILE PARSERS
# ─────────────────────────────────────────────────────────────────────────────

def extract_strings_from_binary(data: bytes, min_length: int = 5) -> list[str]:
    pattern = re.compile(b'[ -~]{' + str(min_length).encode() + b',}')
    results = []
    for match in pattern.finditer(data):
        try:
            s = match.group().decode('ascii', errors='ignore').strip()
            if len(s) >= min_length:
                results.append(s)
        except Exception:
            continue
    return results

def parse_tec19_files(uploaded_files: list) -> list[dict]:
    """Extracts job descriptions and exact completion dates from Word documents."""
    DATE_PATTERN = re.compile(r'\b(\d{2})[/-](\d{2})[/-](\d{2,4})\b')
    all_entries = []

    for f in uploaded_files:
        try:
            file_bytes = f.getvalue()
            
            # Extract raw text lines safely
            if f.name.lower().endswith('.docx'):
                try:
                    doc = Document(io.BytesIO(file_bytes))
                    raw_lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                    if not raw_lines:
                        raise ValueError("Empty docx")
                except Exception:
                    raw_lines = extract_strings_from_binary(file_bytes)
            else:
                raw_lines = extract_strings_from_binary(file_bytes)

            if not raw_lines:
                continue

            current_job_text = None
            collected_dates = []

            for line in raw_lines:
                line = line.strip()
                if not line:
                    continue

                alpha_count = sum(c.isalpha() for c in line)
                dates_found = DATE_PATTERN.findall(line)
                is_date_line = bool(dates_found) and alpha_count < 8

                # Process pure date lines
                if is_date_line:
                    for d_str, m_str, y_str in dates_found:
                        day, month, year = int(d_str), int(m_str), int(y_str)
                        if year < 100:
                            year += 2000
                        if 1 <= day <= 31 and 1 <= month <= 12 and 2010 <= year <= 2099:
                            try:
                                collected_dates.append(datetime(year, month, day))
                            except ValueError:
                                pass
                                
                # Process text/job lines
                elif alpha_count >= 4:
                    # Flush previous job if we have a new one
                    if current_job_text and collected_dates:
                        if len(current_job_text) > 8 and is_action_entry(current_job_text):
                            all_entries.append({
                                'text': current_job_text,
                                'date': collected_dates[-1], # Last date is completion
                                'normalized_text': normalize_maritime_text(current_job_text),
                                'source': f.name
                            })

                    if dates_found:
                        text_part = DATE_PATTERN.sub('', line).strip()
                        for d_str, m_str, y_str in dates_found:
                            day, month, year = int(d_str), int(m_str), int(y_str)
                            if year < 100:
                                year += 2000
                            if 1 <= day <= 31 and 1 <= month <= 12 and 2010 <= year <= 2099:
                                try:
                                    collected_dates = [datetime(year, month, day)]
                                except ValueError:
                                    pass
                                break
                                
                        if len(text_part) > 8:
                            current_job_text = text_part
                            if collected_dates and is_action_entry(text_part):
                                all_entries.append({
                                    'text': current_job_text,
                                    'date': collected_dates[-1],
                                    'normalized_text': normalize_maritime_text(current_job_text),
                                    'source': f.name
                                })
                                current_job_text = None
                                collected_dates = []
                    else:
                        current_job_text = line
                        collected_dates = []

            # Flush the final entry
            if current_job_text and collected_dates:
                if len(current_job_text) > 8 and is_action_entry(current_job_text):
                    all_entries.append({
                        'text': current_job_text,
                        'date': collected_dates[-1],
                        'normalized_text': normalize_maritime_text(current_job_text),
                        'source': f.name
                    })

        except Exception as e:
            st.error(f"❌ Error parsing {f.name}: {e}")

    # Deduplicate exact same entries
    seen = set()
    unique_entries = []
    for entry in all_entries:
        key = (entry['text'][:80], entry['date'])
        if key not in seen:
            seen.add(key)
            unique_entries.append(entry)

    return unique_entries

def parse_pms_excel(file_bytes: bytes) -> tuple[pd.DataFrame | None, str | None]:
    """Extracts PMS components, claimed hours, and monthly historical running hours."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        return None, f"Cannot open Excel file: {e}"

    import warnings
    warnings.filterwarnings('ignore', category=UserWarning)

    target_sheet = wb.active
    for name in wb.sheetnames:
        if 'PMS' in name.upper():
            target_sheet = wb[name]
            break

    df_scan = pd.read_excel(io.BytesIO(file_bytes), sheet_name=target_sheet.title, engine='openpyxl', header=None)

    header_idx = 7
    for idx, row in df_scan.head(15).iterrows():
        row_str = ' '.join(str(v).upper() for v in row if pd.notna(v))
        if ('ITEMS' in row_str or 'CODE' in row_str) and 'INTERVAL' in row_str:
            header_idx = idx
            break

    df = pd.read_excel(io.BytesIO(file_bytes), sheet_name=target_sheet.title, engine='openpyxl', header=header_idx, dtype={0: str})
    df.columns = [str(c).strip().upper() if pd.notna(c) else f'COL_{i}' for i, c in enumerate(df.columns)]

    def find_col(keywords):
        for keyword in keywords:
            for col in df.columns:
                if keyword in col:
                    return col
        return None

    # Map core columns
    cmap = {
        'code': df.columns[0] if len(df.columns) > 0 else None,
        'component': find_col(['ITEMS', 'DESCRIPTION', 'EQUIPMENT']) or (df.columns[1] if len(df.columns) > 1 else None),
        'interval': find_col(['INTERVAL']),
        'date': find_col(['DATE OF LAST', 'LAST INSPECTION', 'LAST OVERHAUL']),
        'hrs_end_yr': find_col(['OPERATING HOURS AT', 'END OF LAST']),
        'current_hrs': find_col(['CURRENT OPERAT'])
    }
    
    # Fallbacks for complex headers
    if not cmap['date']:
        for col in df.columns[4:10]:
            if 'DATE' in col:
                cmap['date'] = col
                break
                
    if not cmap['hrs_end_yr'] and cmap['date']:
        try:
            date_pos = list(df.columns).index(cmap['date'])
            if date_pos + 1 < len(df.columns):
                cmap['hrs_end_yr'] = df.columns[date_pos + 1]
        except ValueError:
            pass
            
    if not cmap['current_hrs'] and cmap['hrs_end_yr']:
        try:
            pos = list(df.columns).index(cmap['hrs_end_yr'])
            if pos + 1 < len(df.columns):
                cmap['current_hrs'] = df.columns[pos + 1]
        except ValueError:
            pass

    # Map monthly tracking columns for Vector 2 Math
    month_keys = ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']
    monthly_cols = {}
    for m in month_keys:
        found = find_col([m.upper()])
        if found:
            monthly_cols[m] = found

    output_rows = []
    for _, row in df.iterrows():
        code = str(row.get(cmap.get('code'), '')).strip()
        comp = str(row.get(cmap.get('component'), '')).strip()

        if not re.match(r'^[A-Z]{2,4}-\d{2}-\d{2,3}$', code):
            continue
        if not comp or comp.upper() in ('NAN', 'NONE', ''):
            continue

        raw_date = row.get(cmap['date']) if cmap.get('date') else None
        last_date = None
        if pd.notna(raw_date):
            if isinstance(raw_date, (datetime, date)):
                last_date = datetime.combine(raw_date, datetime.min.time()) if isinstance(raw_date, date) and not isinstance(raw_date, datetime) else raw_date
            else:
                for fmt in ['%d-%m-%Y', '%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y']:
                    try:
                        last_date = datetime.strptime(str(raw_date)[:10], fmt)
                        break
                    except ValueError:
                        pass

        def safe_float(val):
            try:
                return float(str(val).replace(',', ''))
            except (ValueError, TypeError):
                return 0.0

        interval_raw = row.get(cmap.get('interval'))
        interval_hrs = safe_float(interval_raw) if pd.notna(interval_raw) else None
        
        claimed_raw = row.get(cmap.get('current_hrs'))
        claimed_hrs = safe_float(claimed_raw) if pd.notna(claimed_raw) else None

        row_data = {
            'code': code,
            'component': comp,
            'normalized_component': normalize_maritime_text(comp),
            'interval_hrs': interval_hrs,
            'last_oh_date': last_date,
            'claimed_current_hrs': claimed_hrs,
            'hrs_end_last_yr': safe_float(row.get(cmap.get('hrs_end_yr'))),
        }
        
        for m, c in monthly_cols.items():
            if c:
                row_data[f'hrs_{m}'] = safe_float(row.get(c))
                
        output_rows.append(row_data)

    return pd.DataFrame(output_rows), None

# ─────────────────────────────────────────────────────────────────────────────
# VECTOR 2 & 3: THE TRUTH RECONSTRUCTION ENGINE
# ─────────────────────────────────────────────────────────────────────────────

def calculate_true_hours(pms_row: dict, anchor_date: datetime | None, audit_date: datetime) -> tuple[float | None, str]:
    """Vector 2: Calculates TRUE hours based on TEC exact date and historical monthly logs."""
    if not anchor_date:
        return None, "No Valid TEC Anchor"
    
    m_keys = ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']
    audit_month = min(audit_date.month, 12)
    
    if anchor_date.year == audit_date.year:
        # Sum from the month of the physical overhaul to the audit month
        true_hrs = sum(pms_row.get(f'hrs_{m_keys[i]}', 0.0) for i in range(anchor_date.month - 1, audit_month))
        return true_hrs, "Current Year Math"
        
    elif anchor_date.year < audit_date.year:
        # Use End-of-Year baseline plus this year's running hours
        baseline = pms_row.get('hrs_end_last_yr', 0.0)
        current_yr_total = sum(pms_row.get(f'hrs_{m_keys[i]}', 0.0) for i in range(audit_month))
        return baseline + current_yr_total, "Prior Year Baseline + Math"
        
    return None, "Future Date Error"

def triage_delta(calculated_true: float | None, claimed: float | None, interval: float | None) -> tuple[str, str]:
    """Vector 3: Exposes the exact reality of the crew's claims."""
    if calculated_true is None:
        return "👻 GHOST", "Missing valid TEC-19 proof."
    
    if interval and calculated_true > interval:
        return "🔴 DANGER", f"True hours ({calculated_true:.0f}) exceed {interval:.0f} hr safety interval."

    if claimed is None or claimed == 0.0:
        return "🟡 INCOMPLETE", f"Crew did not log current hours. True hours are {calculated_true:.0f}."

    # Compare True vs Claimed (Allow 5% or 10hr tolerance for pro-rata rounding)
    delta = claimed - calculated_true
    tolerance = max(10, calculated_true * 0.05)

    if abs(delta) <= tolerance:
        return "✅ PERFECT", f"Claimed ({claimed:.0f}) aligns with Truth ({calculated_true:.0f})."
    
    if delta > tolerance:
        return "🟠 FORGOTTEN", f"Crew claims {claimed:.0f}, but Truth is {calculated_true:.0f}. They forgot to reset the counter."
    
    if delta < -tolerance:
        return "🔴 FRAUD", f"Crew claims {claimed:.0f}, but Truth is {calculated_true:.0f}. Hours suspiciously under-reported."
        
    return "UNKNOWN", "Unhandled condition."

# ─────────────────────────────────────────────────────────────────────────────
# VECTOR 1: GLOBAL BIPARTITE VETO ENGINE
# ─────────────────────────────────────────────────────────────────────────────

def run_truth_engine(pms_df: pd.DataFrame, tec_entries: list[dict], audit_date: datetime, threshold: int) -> dict:
    results = {
        'perfect_sync': [], 'forgotten_reset': [], 'fraud_risk': [], 
        'overdue': [], 'ghost': [], 'unlogged': [], 'date_conflicts': []
    }
    
    # 1. Bipartite Matrix Generation
    candidates = []
    
    for p_idx, p_row in pms_df.iterrows():
        p_sys = extract_system_from_code(p_row['code'])
        p_cyl = extract_cylinder_from_code(p_row['code'])
        p_norm = p_row['normalized_component']
        
        for t_idx, t_row in enumerate(tec_entries):
            t_text = t_row['text'].upper()
            t_sys = extract_system_from_text(t_text)
            t_cyl = extract_cylinder_from_text(t_text)
            
            # 🛑 VETO GATES
            if p_sys in FORBIDDEN_KEYWORDS:
                if any(k in t_text for k in FORBIDDEN_KEYWORDS[p_sys]):
                    continue # Hard Quarantine Veto
                    
            if p_sys and t_sys and p_sys != t_sys:
                continue # Hard System Veto
                
            if p_sys in ['MAIN ENGINE', 'DIESEL GENERATOR']:
                if p_cyl and t_cyl and p_cyl != t_cyl:
                    continue # Hard Cylinder Veto
            
            # 🟢 PROTECTED FUZZY MATCH
            score_set = fuzz.token_set_ratio(p_norm, t_row['normalized_text'])
            score_sort = fuzz.token_sort_ratio(p_norm, t_row['normalized_text'])
            score = max(score_set, score_sort * 0.95)
            
            # Anchor Bonus
            if p_cyl and t_cyl and p_cyl == t_cyl:
                score = min(100, score + 15)
            
            if score >= threshold:
                candidates.append((score, p_idx, t_idx))

    # 2. The 1-to-1 Consumption Lock
    candidates.sort(key=lambda x: x[0], reverse=True)
    claimed_pms = set()
    claimed_tec = set()
    final_matches = {}
    
    for score, p, t in candidates:
        if p not in claimed_pms and t not in claimed_tec:
            final_matches[p] = (t, score)
            claimed_pms.add(p)
            claimed_tec.add(t) # Evidence is permanently locked

    # 3. Vector 2 & 3 Integration
    for p_idx, p_row in pms_df.iterrows():
        code = p_row['code']
        comp = p_row['component']
        claimed_hrs = p_row['claimed_current_hrs']
        interval = p_row['interval_hrs']
        p_date = p_row['last_oh_date']
        
        if p_idx not in final_matches:
            results['ghost'].append({
                'Code': code, 'Component': comp, 
                'PMS Date': p_date.strftime('%Y-%m-%d') if pd.notna(p_date) else '—',
                'Claimed Hrs': f"{claimed_hrs:.0f}" if claimed_hrs else '—', 
                'Status': '👻 GHOST OVERHAUL — No Valid TEC Evidence',
            })
            continue

        t_idx, score = final_matches[p_idx]
        t_row = tec_entries[t_idx]
        t_date = t_row['date']

        # Vector 2: Reconstruct Truth
        true_hrs, math_method = calculate_true_hours(p_row.to_dict(), t_date, audit_date)
        
        # Vector 3: Expose Reality
        status_flag, msg = triage_delta(true_hrs, claimed_hrs, interval)

        record = {
            'Code': code, 
            'Component': comp, 
            'TEC Anchor Date': t_date.strftime('%Y-%m-%d'),
            'PMS Claimed Date': p_date.strftime('%Y-%m-%d') if pd.notna(p_date) else '—',
            'TEC Evidence': t_row['text'][:80], 
            'Match': f"{score:.0f}%",
            'Calculated TRUE Hrs': f"{true_hrs:.0f}" if true_hrs is not None else '—',
            'Claimed Excel Hrs': f"{claimed_hrs:.0f}" if claimed_hrs else '—',
            'Interval Limit': f"{interval:.0f}" if interval else '—',
            'Verdict': msg
        }

        # Route to correct bucket based on Vector 3 Expose
        if pd.notna(p_date) and t_date and (t_date.year != p_date.year or t_date.month != p_date.month):
            record['Verdict'] = f"DATE CONFLICT: Work done {t_date.strftime('%b %Y')} but claimed {p_date.strftime('%b %Y')}."
            results['date_conflicts'].append(record)
        elif "DANGER" in status_flag:
            results['overdue'].append(record)
        elif "FRAUD" in status_flag:
            results['fraud_risk'].append(record)
        elif "FORGOTTEN" in status_flag:
            results['forgotten_reset'].append(record)
        elif "PERFECT" in status_flag:
            results['perfect_sync'].append(record)
        else:
            results['ghost'].append(record)

    # 4. Find Unlogged Jobs
    for i, t in enumerate(tec_entries):
        if i not in claimed_tec:
            results['unlogged'].append({
                'TEC Date': t['date'].strftime('%Y-%m-%d'), 
                'Evidence': t['text'][:120], 
                'Status': '📋 UNLOGGED — Performed but not added to PMS Excel'
            })

    return results

# ─────────────────────────────────────────────────────────────────────────────
# REPORT & UI ENGINE
# ─────────────────────────────────────────────────────────────────────────────

def generate_excel_report(results: dict, vessel: str, d: datetime) -> bytes:
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine='xlsxwriter') as w:
        wb = w.book
        f_hdr = wb.add_format({'bold': True, 'bg_color': '#111111', 'font_color': '#FFFFFF', 'border': 1})
        
        sheets = [
            ('🚨 Fraud & Typos', results['fraud_risk'], '#FFE4E1'),
            ('🔴 Overdue Danger', results['overdue'], '#FFCCCC'),
            ('🟠 Forgotten Resets', results['forgotten_reset'], '#FFF3CD'),
            ('⚠️ Date Conflicts', results['date_conflicts'], '#FFFDE7'),
            ('👻 Ghosts', results['ghost'], '#F0F0F0'),
            ('✅ Perfect Syncs', results['perfect_sync'], '#E8F5E9'),
            ('📋 Unlogged', results['unlogged'], '#FFFFFF'),
        ]
        
        # Write Summary
        summ_data = {
            'Metric': [s[0] for s in sheets], 
            'Count': [len(s[1]) for s in sheets]
        }
        pd.DataFrame(summ_data).to_excel(w, sheet_name='Summary', index=False)
        
        # Write Data Sheets
        for name, data, color in sheets:
            safe_name = name[-30:] # Excel sheet name limit
            if not data:
                pd.DataFrame([{'Note': 'No items in this category.'}]).to_excel(w, sheet_name=safe_name, index=False)
                continue
                
            df = pd.DataFrame(data)
            df.to_excel(w, sheet_name=safe_name, index=False)
            ws = w.sheets[safe_name]
            fmt = wb.add_format({'bg_color': color, 'border': 1})
            
            for col_idx, col_name in enumerate(df.columns):
                ws.write(0, col_idx, col_name, f_hdr)
                ws.set_column(col_idx, col_idx, 25)
                
            for r_idx in range(len(df)):
                for c_idx in range(len(df.columns)):
                    ws.write(r_idx + 1, c_idx, str(df.iloc[r_idx, c_idx]), fmt)
                    
    return out.getvalue()

def render_card(val, label, color):
    return f"""
    <div class="metric-card" style="border-top: 4px solid {color}">
        <div class="metric-value" style="color:{color}">{val}</div>
        <div class="metric-label">{label}</div>
    </div>
    """

def main():
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@400;600;700&family=JetBrains+Mono:wght@400;700&display=swap');
    html, body, [class*="css"] {font-family: 'Space Grotesk', sans-serif;}
    .main {background: #0d0d1a; color: white;}
    .metric-card {background: #1a1a2e; border-radius: 8px; padding: 1.5rem; text-align: center; margin-bottom: 1rem; box-shadow: 0 4px 6px rgba(0,0,0,0.3);}
    .metric-value {font-size: 2.5rem; font-weight: 700; font-family: 'JetBrains Mono';}
    .metric-label {font-size: 0.85rem; color: #8888aa; text-transform: uppercase; letter-spacing: 1px; margin-top: 0.5rem;}
    div[data-testid="stSidebar"] {background: #0a0a14; border-right: 1px solid #222233;}
    .stButton>button {background: linear-gradient(135deg, #6c63ff, #5a54d6); color: white; font-weight: bold; border-radius: 8px; border: none;}
    .stButton>button:hover {background: linear-gradient(135deg, #5a54d6, #4843b4); box-shadow: 0 0 15px rgba(108, 99, 255, 0.4);}
    </style>
    """, unsafe_allow_html=True)

    with st.sidebar:
        st.markdown('<h2 style="color:#6c63ff; font-weight:bold; font-family:\'JetBrains Mono\';">⚓ THE TRUTH ENGINE</h2>', unsafe_allow_html=True)
        st.markdown('<p style="color:#88a; font-size:0.9rem; text-transform:uppercase;">Zero-Trust Forensic Audit v3.0</p><hr style="border-color:#334">', unsafe_allow_html=True)
        
        vessel_name = st.text_input("Vessel Name", "MV ALEXIS")
        audit_date_input = st.date_input("Audit Date", date.today())
        match_threshold = st.slider("Fuzzy Match Threshold (%)", min_value=40, max_value=90, value=58, step=2)
        
        st.markdown("<br><b>1. UPLOAD PMS EXCEL</b>", unsafe_allow_html=True)
        pms_file = st.file_uploader("PMS Data", type=['xlsx', 'xls'], label_visibility="collapsed")
        
        st.markdown("<br><b>2. UPLOAD TEC-19 DIARY</b>", unsafe_allow_html=True)
        tec_files = st.file_uploader("TEC-19 Logs", type=['doc', 'docx'], accept_multiple_files=True, label_visibility="collapsed")
        
        st.markdown("<br>", unsafe_allow_html=True)
        run_clicked = st.button("▶ EXECUTE TRUTH AUDIT", type='primary', use_container_width=True, disabled=not (pms_file and tec_files))

    if run_clicked:
        audit_dt = datetime.combine(audit_date_input, datetime.min.time())
        
        with st.status("🔬 Extracting Truth Vectors...", expanded=True) as status:
            st.write("📊 Parsing PMS Ledger...")
            pms_df, pms_error = parse_pms_excel(pms_file.getvalue())
            if pms_error:
                st.error(pms_error)
                status.update(label="Audit Failed", state="error")
                st.stop()
                
            st.write("📝 Parsing TEC-19 Work Diary...")
            tec_entries = parse_tec19_files(tec_files)
            if not tec_entries:
                st.error("No valid maintenance action entries found in TEC-19.")
                status.update(label="Audit Failed", state="error")
                st.stop()
                
            st.write("🔗 Executing Bipartite Veto Match & Math Reconstruction...")
            results = run_truth_engine(pms_df, tec_entries, audit_dt, match_threshold)
            
            status.update(label="✅ Audit Complete.", state="complete")
            
        # Display Results Dashboard
        c1, c2, c3, c4, c5 = st.columns(5)
        c1.markdown(render_card(len(results['fraud_risk']), "Fraud / Typo Risk", "#ff4444"), unsafe_allow_html=True)
        c2.markdown(render_card(len(results['overdue']), "Danger Overdue", "#ff9933"), unsafe_allow_html=True)
        c3.markdown(render_card(len(results['forgotten_reset']), "Forgotten Resets", "#ffcc00"), unsafe_allow_html=True)
        c4.markdown(render_card(len(results['ghost']), "Ghosts", "#aa44ff"), unsafe_allow_html=True)
        c5.markdown(render_card(len(results['perfect_sync']), "Perfect Syncs", "#44ff44"), unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)
        
        report_bytes = generate_excel_report(results, vessel_name, audit_dt)
        st.download_button(
            label="📥 DOWNLOAD FORENSIC TRUTH REPORT (EXCEL)", 
            data=report_bytes, 
            file_name=f"Truth_Report_{vessel_name.replace(' ', '_')}.xlsx", 
            use_container_width=True
        )

        st.markdown("<hr>", unsafe_allow_html=True)

        # Tabs for detailed viewing
        tab_names = [
            f"🚨 Fraud/Typos ({len(results['fraud_risk'])})", 
            f"🔴 Overdue ({len(results['overdue'])})", 
            f"🟠 Forgotten Resets ({len(results['forgotten_reset'])})", 
            f"⚠️ Dates ({len(results['date_conflicts'])})", 
            f"👻 Ghosts ({len(results['ghost'])})", 
            f"✅ Perfect ({len(results['perfect_sync'])})"
        ]
        
        tabs = st.tabs(tab_names)
        
        with tabs[0]: st.dataframe(pd.DataFrame(results['fraud_risk']), use_container_width=True)
        with tabs[1]: st.dataframe(pd.DataFrame(results['overdue']), use_container_width=True)
        with tabs[2]: st.dataframe(pd.DataFrame(results['forgotten_reset']), use_container_width=True)
        with tabs[3]: st.dataframe(pd.DataFrame(results['date_conflicts']), use_container_width=True)
        with tabs[4]: st.dataframe(pd.DataFrame(results['ghost']), use_container_width=True)
        with tabs[5]: st.dataframe(pd.DataFrame(results['perfect_sync']), use_container_width=True)

if __name__ == "__main__":
    main()
